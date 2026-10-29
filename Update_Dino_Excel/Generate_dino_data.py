"""
Update the JW3 Dino Habitat Planner workbook from paleo.gg.

Only these columns are rewritten (matched by header text in row 1, so it works
on sheets that don't have every column, e.g. Aviary has no "Diet"):

    DINO, Diet, Era, Size, Likes, Dislikes

Everything else (habitat percentages, fills, fonts, widths, notes) is untouched.

Usage:
    python Update_dino_sheet.py                      # writes JW3_Dino_Habitat_Planner_updated.xlsx
    python Update_dino_sheet.py --dry-run            # only prints what would change
    python Update_dino_sheet.py --inplace            # overwrite the original (a .bak copy is made first)
    python Update_dino_sheet.py -i other.xlsx -o out.xlsx
"""

import argparse
import shutil
import time
from pathlib import Path

import requests
from openpyxl import load_workbook

# --------------------------------------------------------------------------
# Config (from your script)
# --------------------------------------------------------------------------

DEFAULT_WORKBOOK = "JW3_Dino_Habitat_Planner.xlsx"

# NOTE: the long id in this URL is the Next.js "buildId". paleo.gg changes it
# whenever they redeploy the site, at which point every request will 404.
# If that happens, grab the new URL from your browser's network tab.
dino_data_page = "https://www.paleo.gg/_next/data/CvQgqvPTCVwxd4R3DacZM/en/games/jurassic-world-evolution-3/dino-db/[DINONAME].json?dino=[DINONAME]"

REQUEST_DELAY = 0.25   # seconds between real HTTP requests, to be polite to the site
REQUEST_TIMEOUT = 20

ERA_KEYS = {
    "era_earlycretaceous": "Early Cretaceous",
    "era_earlyjurassic": "Early Jurassic",
    "era_earlypermian": "Early Permian",
    "era_holocene": "Holocene",
    "era_latecretaceous": "Late Cretaceous",
    "era_middlejurassic": "Middle Jurassic",
    "era_latejurassic": "Late Jurassic",
    "era_latepermian": "Late Permian",
    "era_latetriassic": "Late Triassic",
    "era_latedevonian": "Late Devonian",
    "era_miocene": "Miocene",
}

DIET_KEYS = {
    "diet_carnivore": "Carnivore",
    "diet_herbivore": "Herbivore",
    "diet_omnivore": "Omnivore",
    "diet_omnivorepiscivore": "Omnivore Piscivore",
    "diet_piscivore": "Piscivore",
    "diet_tallherbivore": "Tall Herbivore",
    "diet_fish": "Piscivore",
    "diet_shoals": "Shoals",
    "diet_shark": "Shark",
}

SIZE_KEYS = {
    "small": "Small",
    "medium": "Medium",
    "large": "Large",
}

# Header text in row 1  ->  key in the dict returned by build_dino_data()
UPDATE_COLUMNS = {
    "dino": "name",
    "diet": "diet",
    "era": "era",
    "size": "size",
    "likes": "likes",
    "dislikes": "dislikes",
}

# --------------------------------------------------------------------------
# Fetching
# --------------------------------------------------------------------------

_json_cache = {}    # slug -> json dict, or None if the dino page doesn't exist
_label_cache = {}   # slug -> display name


def build_dino_url(dino_name: str) -> str:
    return dino_data_page.replace("[DINONAME]", dino_name)


def fetch_dino_json(slug: str):
    """Return the page JSON for a slug, or None if paleo.gg has no such dino (404)."""
    if slug in _json_cache:
        return _json_cache[slug]

    time.sleep(REQUEST_DELAY)
    response = requests.get(build_dino_url(slug), timeout=REQUEST_TIMEOUT)
    if response.status_code == 404:
        _json_cache[slug] = None
        return None
    response.raise_for_status()

    data = response.json()
    _json_cache[slug] = data
    return data


def slug_candidates(sheet_name: str):
    """
    The sheet says 'Tyrannosaurus rex' / 'Indominus rex' etc., while the site
    wants a slug. We don't know the exact slug style for multi-word names, so
    try the likely ones in order until one isn't a 404.
    """
    clean = sheet_name.replace("\u00ad", "").strip().lower()   # strips the soft hyphen in 'Hypsiloph­odon'
    options = [
        clean.replace(" ", ""),    # acrocanthosaurus / tyrannosaurusrex
        clean.replace(" ", "-"),   # tyrannosaurus-rex
        clean.replace(" ", "_"),   # tyrannosaurus_rex
    ]
    seen, result = set(), []
    for opt in options:
        if opt not in seen:
            seen.add(opt)
            result.append(opt)
    return result


def get_name_label(slug: str) -> str:
    """Display name for a slug (used for the Likes / Dislikes lists)."""
    if slug in _label_cache:
        return _label_cache[slug]

    label = ""
    try:
        data = fetch_dino_json(slug)
        if data:
            label = data.get("pageProps", {}).get("detail", {}).get("name", "")
    except requests.RequestException:
        pass

    # Fallback if the entry isn't a dino page (e.g. a group like "scavenger")
    if not label:
        label = slug.replace("-", " ").replace("_", " ").title()

    _label_cache[slug] = label
    return label


def build_dino_data(dino_data: dict) -> dict:
    details = dino_data.get("pageProps", {}).get("detail", {})

    name = details.get("name", "")
    diet = DIET_KEYS.get(details.get("diet", ""), details.get("diet", ""))
    era = ERA_KEYS.get(details.get("era", ""), details.get("era", ""))
    size = SIZE_KEYS.get(details.get("size", ""), details.get("size", ""))

    def labels(key):
        out = []
        for entry in details.get(key, []):
            slug = entry[0] if isinstance(entry, (list, tuple)) else entry
            out.append(get_name_label(slug))
        return out

    return {
        "name": name,
        "diet": diet,
        "era": era,
        "size": size,
        "likes": labels("cohab_like"),
        "dislikes": labels("cohab_dislike"),
    }


def lookup_dino(sheet_name: str):
    """Try each slug candidate; return built data or None if the dino isn't found."""
    for slug in slug_candidates(sheet_name):
        data = fetch_dino_json(slug)
        if data:
            return build_dino_data(data)
    return None


# --------------------------------------------------------------------------
# Workbook updating
# --------------------------------------------------------------------------

def header_map(ws) -> dict:
    """{'dino': 1, 'diet': 2, ...} for the columns we're allowed to update."""
    cols = {}
    for cell in ws[1]:
        if isinstance(cell.value, str):
            key = cell.value.strip().lower()
            if key in UPDATE_COLUMNS:
                cols[key] = cell.column
    return cols


def to_cell_value(value):
    """Lists become one-entry-per-line (matches the sheet's existing style); empty -> blank."""
    if isinstance(value, list):
        value = "\n".join(value)
    return value if value else None


def update_workbook(src: Path, dst: Path, dry_run: bool):
    wb = load_workbook(src)
    changed_cells = 0
    not_found = []
    errors = []

    for ws in wb.worksheets:
        cols = header_map(ws)
        if "dino" not in cols:
            print(f"[{ws.title}] no 'DINO' column, skipping sheet")
            continue

        print(f"\n=== {ws.title}  (updating: {', '.join(k for k in cols)}) ===")

        for row in range(2, ws.max_row + 1):
            sheet_name = ws.cell(row, cols["dino"]).value
            if not sheet_name or not str(sheet_name).strip():
                continue
            sheet_name = str(sheet_name)

            try:
                new = lookup_dino(sheet_name)
            except requests.RequestException as exc:
                errors.append((ws.title, sheet_name, str(exc)))
                print(f"  ! {sheet_name}: request failed ({exc})")
                continue

            if new is None:
                not_found.append((ws.title, sheet_name))
                print(f"  ? {sheet_name}: not found on paleo.gg, row left unchanged")
                continue

            for header, key in UPDATE_COLUMNS.items():
                if header not in cols:
                    continue
                cell = ws.cell(row, cols[header])
                new_value = to_cell_value(new[key])
                if (cell.value or None) != new_value:
                    old_txt = str(cell.value).replace("\n", " | ") if cell.value else "(blank)"
                    new_txt = str(new_value).replace("\n", " | ") if new_value else "(blank)"
                    print(f"  {sheet_name} [{header}]: {old_txt}  ->  {new_txt}")
                    changed_cells += 1
                    if not dry_run:
                        cell.value = new_value   # value only: fills, fonts, wrap are kept

    print("\n" + "-" * 60)
    print(f"Cells {'that would change' if dry_run else 'changed'}: {changed_cells}")
    if not_found:
        print(f"Not found on paleo.gg ({len(not_found)}):")
        for sheet, name in not_found:
            print(f"  - {sheet}: {name}")
    if errors:
        print(f"Request errors ({len(errors)}):")
        for sheet, name, msg in errors:
            print(f"  - {sheet}: {name}: {msg}")

    if dry_run:
        print("Dry run: nothing saved.")
    else:
        wb.save(dst)
        print(f"Saved: {dst}")


def main():
    parser = argparse.ArgumentParser(description="Update the dino habitat workbook from paleo.gg")
    parser.add_argument("-i", "--input", default=DEFAULT_WORKBOOK, help="workbook to read")
    parser.add_argument("-o", "--output", help="where to save (default: <input>_updated.xlsx)")
    parser.add_argument("--inplace", action="store_true", help="overwrite the input file (makes a .bak copy first)")
    parser.add_argument("--dry-run", action="store_true", help="show changes without saving")
    args = parser.parse_args()

    src = Path(args.input)
    if not src.exists():
        raise SystemExit(f"Workbook not found: {src}")

    if args.inplace:
        dst = src
        if not args.dry_run:
            shutil.copy2(src, src.with_suffix(".xlsx.bak"))
    else:
        dst = Path(args.output) if args.output else src.with_name(f"{src.stem}_updated{src.suffix}")

    update_workbook(src, dst, args.dry_run)

if __name__ == "__main__":
    main()